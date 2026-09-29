import sys
import io
import requests
import openpyxl
import pandas as pd
import streamlit as st

# ==============================================================================
# AUTORUNNER: FUNCIONA TANTO LOCALMENTE (DOBLE CLIC) COMO EN STREAMLIT CLOUD
# ==============================================================================
if __name__ == "__main__":
    if not st.runtime.exists():
        from streamlit.web import cli as stcli
        sys.argv = ["streamlit", "run", sys.argv[0]]
        sys.exit(stcli.main())

# ==============================================================================
# CONFIGURACIÓN DE PÁGINA
# ==============================================================================
st.set_page_config(
    page_title="Ahorro Programado Jardín 2026-2027",
    page_icon="💰",
    layout="wide"
)

DRIVE_FILE_ID = "1s8H4eqdO3LyIFwPdjGZG0TwHCRZmtnsw"

# ==============================================================================
# CARGAR DESDE GOOGLE DRIVE (CON CACHÉ REUTILIZABLE)
# ==============================================================================
@st.cache_data(ttl=300)
def cargar_excel_drive(file_id):
    url = f"https://drive.google.com/uc?id={file_id}&export=download"
    session = requests.Session()
    response = session.get(url)
    
    if response.status_code != 200:
        st.error("No se pudo descargar el archivo desde Google Drive.")
        return None

    excel_bytes = io.BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_bytes, data_only=True)
    return wb

def obtener_valor(sheet, celda_principal, celda_secundaria=None):
    val = sheet[celda_principal].value
    if val is None and celda_secundaria:
        val = sheet[celda_secundaria].value
    return val

# ==============================================================================
# ENCABEZADO PRINCIPAL Y RECARGA
# ==============================================================================
col_title, col_btn = st.columns([4, 1])
with col_title:
    st.title("💰 Ahorro Programado Jardín 2026-2027")
    st.caption("Sincronizado en tiempo real desde Google Drive")

with col_btn:
    if st.button("🔄 Recargar Excel", use_container_width=True):
        st.cache_data.clear()
        st.rerun()

st.divider()

wb = cargar_excel_drive(DRIVE_FILE_ID)

if wb is not None:
    sheet = wb.active

    # ==============================================================================
    # 1. FICHA TÉCNICA CONDICIONES DEL AHORRO (E1:J9)
    # ==============================================================================
    tna = obtener_valor(sheet, "H1", "F1")
    tasa_diaria = obtener_valor(sheet, "J2", "F2")
    fecha_inicio = obtener_valor(sheet, "H3", "F3")
    plazo_meses = obtener_valor(sheet, "H4", "F4")
    fecha_fin = obtener_valor(sheet, "H5", "F5")
    total_dias = obtener_valor(sheet, "H6", "F6")
    valor_aporte = obtener_valor(sheet, "H7", "F7")
    dia_aporte = obtener_valor(sheet, "H8", "F8")
    meses_real = sheet["H9"].value
    dias_real = sheet["J9"].value

    tna_val = float(tna) * 100 if isinstance(tna, (int, float)) and tna <= 1 else float(tna or 0)
    fecha_ini_str = fecha_inicio.strftime('%d/%m/%Y') if hasattr(fecha_inicio, 'strftime') else str(fecha_inicio or "-")
    fecha_fin_str = fecha_fin.strftime('%d/%m/%Y') if hasattr(fecha_fin, 'strftime') else str(fecha_fin or "-")
    val_diaria_str = f"{float(tasa_diaria):.7f}".replace(".", ",") if isinstance(tasa_diaria, (int, float)) else str(tasa_diaria)
    val_aporte_str = f"{float(valor_aporte):,.2f}".replace(".", ",") if isinstance(valor_aporte, (int, float)) else str(valor_aporte)

    st.subheader("📋 Condiciones del Ahorro Programado")

    html_condiciones = f"""<style>
.excel-tbl-card {{
    background-color: #ffffff;
    padding: 10px;
    border-radius: 6px;
    box-shadow: 0 3px 10px rgba(0, 0, 0, 0.2);
    display: inline-block;
    margin-bottom: 15px;
}}
.excel-tbl {{
    border-collapse: collapse;
    font-family: Calibri, 'Segoe UI', Arial, sans-serif;
    font-size: 13px;
    color: #000000;
    width: auto;
}}
.excel-tbl td {{
    border: 1px solid #000000;
    padding: 4px 8px;
    vertical-align: middle;
}}
.lbl-yellow {{
    background-color: #FFE600;
    font-weight: bold;
    text-align: right;
    white-space: nowrap;
}}
.val-green {{
    background-color: #E2EFDA;
    text-align: center;
    font-weight: 600;
}}
.val-formula {{
    background-color: #D9E1F2;
    text-align: center;
    font-size: 12px;
}}
.val-tan {{
    background-color: #FFF2CC;
    text-align: center;
    font-weight: 600;
}}
</style>
<div class="excel-tbl-card">
<table class="excel-tbl">
<tr>
<td class="lbl-yellow">Tasa Nominal Anual</td>
<td colspan="4" class="val-green">{tna_val:.2f}%</td>
</tr>
<tr>
<td class="lbl-yellow">Tasa Diaria</td>
<td colspan="3" class="val-formula">((TNA/100)/360)</td>
<td class="val-green" style="font-weight: bold;">{val_diaria_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Fecha inicio</td>
<td colspan="4" class="val-green">{fecha_ini_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Plazo meses</td>
<td colspan="4" class="val-green">{plazo_meses}</td>
</tr>
<tr>
<td class="lbl-yellow">Fecha Fin</td>
<td colspan="4" class="val-tan">{fecha_fin_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Total de Días</td>
<td colspan="4" class="val-tan">{total_dias}</td>
</tr>
<tr>
<td class="lbl-yellow">Valor Aporte Mensual</td>
<td colspan="4" class="val-green">${val_aporte_str}</td>
</tr>
<tr>
<td class="lbl-yellow">Dia de Aporte Mensual</td>
<td colspan="4" class="val-green">{dia_aporte}</td>
</tr>
<tr>
<td class="lbl-yellow">Cantidad de Meses Real</td>
<td class="val-tan" style="width: 15%;">Meses</td>
<td class="val-tan" style="width: 15%;">{meses_real}</td>
<td class="val-tan" style="width: 15%;">Días</td>
<td class="val-tan" style="width: 15%;">{dias_real}</td>
</tr>
</table>
</div>"""

    st.markdown(html_condiciones, unsafe_allow_html=True)

    st.divider()

    # ==============================================================================
    # 2. TABLA RESUMEN DE MOVIMIENTOS DINÁMICA (AZ:BF)
    # ==============================================================================
    st.subheader("📑 Tabla Resumen de Movimientos")

    movements = []
    r = 3  # Fila de inicio (AZ3:BF3)
    
    while True:
        date_val = sheet.cell(row=r, column=52).value  # Col AZ (Fechas)
        if date_val is None or str(date_val).strip() == "":
            break

        dep_m = sheet.cell(row=r, column=53).value or 0     # Col BA
        dep_p = sheet.cell(row=r, column=54).value or 0     # Col BB
        saldo_acum = sheet.cell(row=r, column=55).value or 0 # Col BC
        int_diario = sheet.cell(row=r, column=56).value or 0 # Col BD
        inc_int_d = sheet.cell(row=r, column=57).value or 0  # Col BE
        int_m = sheet.cell(row=r, column=58).value          # Col BF

        is_int_m = int_m is not None and isinstance(int_m, (int, float)) and float(int_m) > 0
        is_dep_p = isinstance(dep_p, (int, float)) and float(dep_p) > 0
        is_dep_m = isinstance(dep_m, (int, float)) and float(dep_m) > 0

        if is_int_m:
            bg_color = "#B4C6E7"  # Azul suave
            tipo_cat = "🟦 Intereses Ganados"
        elif is_dep_p:
            bg_color = "#FFFF00"  # Amarillo
            tipo_cat = "🟨 Aportes Personales"
        elif is_dep_m:
            bg_color = "#FCE4D6"  # Naranja/Rosado
            tipo_cat = "🟧 Aportes Mensuales"
        else:
            bg_color = "#FFFFFF"
            tipo_cat = "Otros"

        movements.append({
            'fecha': date_val,
            'dep_m': float(dep_m) if isinstance(dep_m, (int, float)) else 0.0,
            'dep_p': float(dep_p) if isinstance(dep_p, (int, float)) else 0.0,
            'saldo_acum': float(saldo_acum) if isinstance(saldo_acum, (int, float)) else 0.0,
            'int_diario': float(int_diario) if isinstance(int_diario, (int, float)) else 0.0,
            'inc_int_d': float(inc_int_d) if isinstance(inc_int_d, (int, float)) else 0.0,
            'int_m': float(int_m) if is_int_m else None,
            'bg_color': bg_color,
            'tipo_cat': tipo_cat
        })
        r += 1

    # ==============================================================================
    # CONTROLES DE FILTRO INTERACTIVO (Pills táctiles)
    # ==============================================================================
    opciones_disponibles = ["🟨 Aportes Personales", "🟧 Aportes Mensuales", "🟦 Intereses Ganados"]
    
    st.write("**Filtra los movimientos por categoría:**")
    
    if hasattr(st, "pills"):
        categorias_seleccionadas = st.pills(
            label="Categorías",
            options=opciones_disponibles,
            default=opciones_disponibles,
            selection_mode="multi",
            label_visibility="collapsed"
        )
    else:
        categorias_seleccionadas = st.multiselect(
            label="Selecciona Categorías a mostrar",
            options=opciones_disponibles,
            default=opciones_disponibles,
            label_visibility="collapsed"
        )

    movimientos_filtrados = [
        m for m in movements if m['tipo_cat'] in (categorias_seleccionadas or [])
    ]

    # ==============================================================================
    # TABLA HTML OPTIMIZADA (CABECERA FIJA + ANCHO AJUSTADO + CENTRADO)
    # ==============================================================================
    rows_html = []
    for m in movimientos_filtrados:
        dt_str = m['fecha'].strftime('%d/%m/%Y') if hasattr(m['fecha'], 'strftime') else str(m['fecha'])
        dep_m_str = f"{m['dep_m']:,.2f}".replace(".", ",") if m['dep_m'] > 0 else "0,00"
        dep_p_str = f"{m['dep_p']:,.2f}".replace(".", ",") if m['dep_p'] > 0 else "0,00"
        saldo_str = f"{m['saldo_acum']:,.2f}".replace(".", ",")
        int_d_str = f"{m['int_diario']:.2f}".replace(".", ",")
        inc_d_str = f"{m['inc_int_d']:.2f}".replace(".", ",") if m['inc_int_d'] > 0 else ""
        int_m_str = f"<b>{m['int_m']:,.2f}</b>".replace(".", ",") if m['int_m'] is not None else ""

        rows_html.append(f"""<tr style="background-color: {m['bg_color']};">
<td><b>{dt_str}</b></td>
<td>{dep_m_str}</td>
<td>{dep_p_str}</td>
<td style="font-weight: 600;">{saldo_str}</td>
<td>{int_d_str}</td>
<td>{inc_d_str}</td>
<td>{int_m_str}</td>
</tr>""")

    tabla_movs_body = "\n".join(rows_html)

    html_movs_completo = f"""<style>
.tbl-scroll-wrapper {{
    max-height: 480px;
    max-width: 100%;
    overflow-y: auto;
    overflow-x: auto;
    border: 1px solid #000000;
    border-radius: 6px;
    background-color: #ffffff;
    box-shadow: 0 4px 12px rgba(0,0,0,0.2);
    margin-top: 10px;
    margin-bottom: 25px;
    display: inline-block;
}}
.tbl-sticky-movs {{
    width: auto;
    border-collapse: collapse;
    font-family: Calibri, 'Segoe UI', Arial, sans-serif;
    font-size: 12px;
    color: #000000;
}}
.tbl-sticky-movs th {{
    position: sticky;
    top: 0;
    z-index: 10;
    background-color: #1F2937;
    color: #ffffff;
    border: 1px solid #000000;
    padding: 5px 6px;
    text-align: center;
    font-weight: bold;
    line-height: 1.15;
    white-space: nowrap;
}}
.tbl-sticky-movs td {{
    border: 1px solid #000000;
    padding: 4px 6px;
    vertical-align: middle;
    text-align: center !important;
    white-space: nowrap;
}}
</style>
<div class="tbl-scroll-wrapper">
<table class="tbl-sticky-movs">
<thead>
<tr>
<th>Fechas</th>
<th>Dep.<br>Mensual</th>
<th>Dep.<br>Personales</th>
<th>Saldo<br>Acumulado</th>
<th>Interés<br>Diario</th>
<th>Incr. Int. Diario<br>por Aportes</th>
<th>Interés<br>Mensual</th>
</tr>
</thead>
<tbody>
{tabla_movs_body if rows_html else '<tr><td colspan="7" style="text-align:center; padding:15px; font-weight:bold;">No hay movimientos para los filtros seleccionados</td></tr>'}
</tbody>
</table>
</div>"""

    st.markdown(html_movs_completo, unsafe_allow_html=True)